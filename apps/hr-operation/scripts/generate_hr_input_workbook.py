#!/usr/bin/env python3
"""Generate the macro-free HR input wizard used by the SharePoint pipeline pilot."""

from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter


NAVY = "172554"
INDIGO = "4F46E5"
INDIGO_LIGHT = "EEF2FF"
SLATE = "475569"
SLATE_LIGHT = "F1F5F9"
WHITE = "FFFFFF"
GREEN = "047857"
GREEN_LIGHT = "D1FAE5"
AMBER = "B45309"
AMBER_LIGHT = "FEF3C7"
RED = "B91C1C"
RED_LIGHT = "FEE2E2"
INPUT_FILL = "FFFDF2"
FORMULA_FILL = "F8FAFC"
GRID = "CBD5E1"

MAX_ROWS = 500
HEADER_ROW = 5
FIRST_DATA_ROW = HEADER_ROW + 1
LAST_DATA_ROW = FIRST_DATA_ROW + MAX_ROWS - 1


def clean(value: object) -> str:
    return str(value or "").replace("\n", " ").strip()


def unique(values: Iterable[object]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        text = clean(value)
        key = text.casefold()
        if text and key not in seen:
            seen.add(key)
            result.append(text)
    return result


def read_reference_lists(job_file: Path | None, employee_file: Path | None) -> dict[str, list[str] | list[tuple[str, str]]]:
    result: dict[str, list[str] | list[tuple[str, str]]] = {
        "departments": [],
        "teams": [],
        "branches": [],
        "grades": [],
        "levels": [],
        "jobs": [],
    }

    if job_file and job_file.exists():
        wb = load_workbook(job_file, read_only=True, data_only=True)
        if "Database" in wb.sheetnames:
            ws = wb["Database"]
            rows = list(ws.iter_rows(min_row=4, values_only=True))
            result["departments"] = unique(row[13] if len(row) > 13 else None for row in rows)
            result["teams"] = unique(row[11] if len(row) > 11 else None for row in rows)
            result["branches"] = unique(row[6] if len(row) > 6 else None for row in rows)
            result["grades"] = unique(row[5] if len(row) > 5 else None for row in rows)
            result["levels"] = unique(row[9] if len(row) > 9 else None for row in rows)

        jobs: list[tuple[str, str]] = []
        seen_jobs: set[str] = set()
        if "Job Titles" in wb.sheetnames:
            ws = wb["Job Titles"]
            for row in ws.iter_rows(min_row=2, values_only=True):
                vi = clean(row[1] if len(row) > 1 else None)
                en = clean(row[2] if len(row) > 2 else None)
                if vi and vi.casefold() not in seen_jobs:
                    seen_jobs.add(vi.casefold())
                    jobs.append((vi, en))
        result["jobs"] = jobs

    if employee_file and employee_file.exists():
        wb = load_workbook(employee_file, read_only=True, data_only=True)
        if "MIDOMAX 2026" in wb.sheetnames:
            ws = wb["MIDOMAX 2026"]
            source_branches = [row[7].value for row in ws.iter_rows(min_row=7) if len(row) > 7]
            result["branches"] = unique([*(result["branches"] or []), *source_branches])

    fallbacks = {
        "departments": ["Ban giám đốc", "Phòng Kinh doanh", "Phòng Marketing", "Phòng NSHC", "Phòng TCKT", "Phòng cung ứng", "Phòng IT"],
        "teams": ["Quản lý phòng", "Nhân sự", "Kế toán", "Mua hàng", "Kho", "Bán lẻ", "Ecom"],
        "branches": ["HN", "HCM"],
        "grades": ["1", "2", "3", "4", "5", "6"],
        "levels": ["1.1", "1.2", "2.1", "2.2", "3.1", "3.2", "4.1", "4.2", "5.1", "5.2", "6.1", "6.2"],
    }
    for key, fallback in fallbacks.items():
        if not result[key]:
            result[key] = fallback
    if not result["jobs"]:
        result["jobs"] = [("Nhân viên", "Executive"), ("Chuyên viên", "Specialist"), ("Trưởng nhóm", "Team Lead"), ("Trưởng phòng", "Manager")]
    return result


def title(ws, text: str, subtitle: str) -> None:
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:Q1")
    ws["A1"] = text
    ws["A1"].font = Font(size=20, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells("A2:Q2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=10, color=SLATE)
    ws["A2"].alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 24


def style_headers(ws, headers: list[str]) -> None:
    thin = Side(style="thin", color=GRID)
    for col, header in enumerate(headers, 1):
        cell = ws.cell(HEADER_ROW, col, header)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(bold=True, color=WHITE, size=9)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.row_dimensions[HEADER_ROW].height = 42


def add_table(ws, name: str, columns: int) -> None:
    ref = f"A{HEADER_ROW}:{ws.cell(LAST_DATA_ROW, columns).coordinate}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def add_list_validation(ws, column: str, named_range: str, prompt: str) -> None:
    validation = DataValidation(type="list", formula1=f"={named_range}", allow_blank=True)
    validation.errorTitle = "Giá trị không hợp lệ"
    validation.error = "Hãy chọn một giá trị trong danh sách."
    validation.promptTitle = "Chọn từ danh mục"
    validation.prompt = prompt
    validation.showErrorMessage = True
    validation.showInputMessage = True
    validation.errorStyle = "stop"
    ws.add_data_validation(validation)
    validation.add(f"{column}{FIRST_DATA_ROW}:{column}{LAST_DATA_ROW}")


def add_custom_validation(ws, cell_range: str, formula: str, title_text: str, message: str, prompt: str) -> None:
    validation = DataValidation(type="custom", formula1=formula, allow_blank=True)
    validation.errorTitle = title_text
    validation.error = message
    validation.promptTitle = "Quy tắc nhập"
    validation.prompt = prompt
    validation.showErrorMessage = True
    validation.showInputMessage = True
    validation.errorStyle = "stop"
    ws.add_data_validation(validation)
    validation.add(cell_range)


def style_input_sheet(ws, formula_columns: set[int], widths: list[float]) -> None:
    ws.freeze_panes = f"C{FIRST_DATA_ROW}"
    ws.auto_filter.ref = f"A{HEADER_ROW}:{ws.cell(LAST_DATA_ROW, len(widths)).coordinate}"
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=FIRST_DATA_ROW, max_row=LAST_DATA_ROW, max_col=len(widths)):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in formula_columns)
            cell.protection = Protection(locked=cell.column in formula_columns)
            cell.fill = PatternFill("solid", fgColor=FORMULA_FILL if cell.column in formula_columns else INPUT_FILL)
            cell.font = Font(size=9, color=SLATE)
    ws.auto_filter.ref = f"A{HEADER_ROW}:{ws.cell(LAST_DATA_ROW, len(widths)).coordinate}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.protection.sheet = True
    ws.protection.password = "hr-template"
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False


def add_row_quality_formatting(ws, last_col_letter: str) -> None:
    data_range = f"A{FIRST_DATA_ROW}:{last_col_letter}{LAST_DATA_ROW}"
    ws.conditional_formatting.add(
        data_range,
        FormulaRule(formula=[f'$A{FIRST_DATA_ROW}="LỖI"'], fill=PatternFill("solid", fgColor=RED_LIGHT)),
    )
    ws.conditional_formatting.add(
        data_range,
        FormulaRule(formula=[f'$A{FIRST_DATA_ROW}="HỢP LỆ"'], fill=PatternFill("solid", fgColor=GREEN_LIGHT)),
    )


def build_start_sheet(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "BAT_DAU"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    for column in "BCDEFGHIJKLMNO":
        ws.column_dimensions[column].width = 12
    ws.merge_cells("B2:O3")
    ws["B2"] = "HR DATA INPUT WIZARD"
    ws["B2"].fill = PatternFill("solid", fgColor=NAVY)
    ws["B2"].font = Font(size=24, bold=True, color=WHITE)
    ws["B2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("B4:O4")
    ws["B4"] = "Form chuẩn cho pipeline SharePoint ↔ HR App · Phiên bản 1.0"
    ws["B4"].font = Font(size=11, color=SLATE)
    ws["B4"].alignment = Alignment(horizontal="center")

    cards = [
        (2, 5, "Bước 1", "Nhập hồ sơ nhân sự", "Mã, họ tên, liên hệ và thông tin cơ bản", "#'01_HO_SO'!A1"),
        (7, 10, "Bước 2", "Nhập thông tin công việc", "Phòng ban, chức danh, quản lý và lương", "#'02_CONG_VIEC'!A1"),
        (12, 15, "Bước 3", "Kiểm tra chất lượng", "Sửa toàn bộ lỗi trước khi đưa lên SharePoint", "#'03_KIEM_TRA'!A1"),
    ]
    for start_col, end_col, step, label, description, link in cards:
        for row in range(6, 11):
            for col in range(start_col, end_col + 1):
                cell = ws.cell(row, col)
                cell.fill = PatternFill("solid", fgColor=INDIGO_LIGHT)
                cell.border = Border(
                    left=Side(style="medium" if col == start_col else None, color=INDIGO),
                    right=Side(style="medium" if col == end_col else None, color=INDIGO),
                    top=Side(style="medium" if row == 6 else None, color=INDIGO),
                    bottom=Side(style="medium" if row == 10 else None, color=INDIGO),
                )
        ws.merge_cells(start_row=6, start_column=start_col, end_row=6, end_column=end_col)
        ws.merge_cells(start_row=7, start_column=start_col, end_row=7, end_column=end_col)
        ws.merge_cells(start_row=8, start_column=start_col, end_row=10, end_column=end_col)
        step_cell = ws.cell(6, start_col)
        title_cell = ws.cell(7, start_col)
        desc_cell = ws.cell(8, start_col)
        step_cell.value = step
        step_cell.font = Font(bold=True, color=INDIGO, size=10)
        title_cell.value = label
        title_cell.font = Font(bold=True, color=NAVY, size=13)
        title_cell.hyperlink = link
        title_cell.style = "Hyperlink"
        desc_cell.value = description
        desc_cell.font = Font(color=SLATE, size=9)
        desc_cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.merge_cells("B13:O13")
    ws["B13"] = "QUY TẮC QUAN TRỌNG"
    ws["B13"].font = Font(bold=True, color=WHITE)
    ws["B13"].fill = PatternFill("solid", fgColor=INDIGO)
    rules = [
        "1. Giữ Mã nhân viên dưới dạng text; mã 00001 không được đổi thành 1.",
        "2. Chỉ nhập vào ô màu vàng nhạt. Ô xám là công thức kiểm tra và đã được khóa.",
        "3. Dùng danh sách chọn cho trạng thái, giới tính, phòng ban, bộ phận, chức danh và chi nhánh.",
        "4. Không xóa/đổi tên sheet, tiêu đề cột hoặc Excel Table; pipeline dùng các định danh này.",
        "5. Chỉ đưa file lên SharePoint khi trang 03_KIEM_TRA hiển thị SẴN SÀNG.",
    ]
    for offset, rule in enumerate(rules, 14):
        ws.merge_cells(start_row=offset, start_column=2, end_row=offset, end_column=15)
        ws.cell(offset, 2, rule)
        ws.cell(offset, 2).font = Font(size=10, color=SLATE)
        ws.cell(offset, 2).alignment = Alignment(wrap_text=True)

    ws.merge_cells("B21:O21")
    ws["B21"] = "Danh mục có thể chỉnh tại sheet DANH_MUC. Các công thức kiểm tra chỉ nhằm đảm bảo chất lượng nhập liệu; quyền truy cập dữ liệu nhạy cảm vẫn do HR App và RLS kiểm soát."
    ws["B21"].font = Font(italic=True, color=AMBER, size=9)
    ws["B21"].alignment = Alignment(wrap_text=True)
    ws.freeze_panes = "B6"


def build_profile_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("01_HO_SO")
    title(ws, "BƯỚC 1 · HỒ SƠ NHÂN SỰ", "Nhập một dòng cho mỗi nhân viên. Ô vàng là vùng nhập; cột trạng thái/lỗi được tính tự động.")
    headers = [
        "Trạng thái kiểm tra", "Lỗi dữ liệu", "Mã nhân viên", "Họ và tên", "Ngày sinh", "Giới tính",
        "Email cá nhân", "Email công ty", "Số điện thoại", "Địa chỉ hiện tại", "Ghi chú",
    ]
    style_headers(ws, headers)
    widths = [18, 44, 16, 28, 14, 12, 28, 28, 17, 42, 36]
    style_input_sheet(ws, {1, 2}, widths)
    for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        ws.cell(row, 1, f'=IF($C{row}="","CHƯA NHẬP",IF($B{row}="","HỢP LỆ","LỖI"))')
        ws.cell(row, 2, (
            f'=IF($C{row}="","",TEXTJOIN(" | ",TRUE,'
            f'IF($C{row}="","Thiếu mã nhân viên",""),'
            f'IF(COUNTIF($C${FIRST_DATA_ROW}:$C${LAST_DATA_ROW},$C{row})>1,"Trùng mã nhân viên",""),'
            f'IF($D{row}="","Thiếu họ tên",""),'
            f'IF(AND($E{row}<>"",OR(NOT(ISNUMBER($E{row})),$E{row}<DATE(1940,1,1),$E{row}>TODAY())),"Ngày sinh không hợp lệ",""),'
            f'IF(AND($G{row}<>"",OR(ISERROR(SEARCH("@",$G{row})),ISERROR(SEARCH(".",$G{row})))),"Email cá nhân không hợp lệ",""),'
            f'IF(AND($H{row}<>"",OR(ISERROR(SEARCH("@",$H{row})),ISERROR(SEARCH(".",$H{row})))),"Email công ty không hợp lệ",""),'
            f'IF(AND($I{row}<>"",OR(LEN($I{row})<9,LEN($I{row})>15,NOT(ISNUMBER(--$I{row})))),"Số điện thoại không hợp lệ","")))'
        ))
        ws.cell(row, 3).number_format = "@"
        ws.cell(row, 5).number_format = "dd/mm/yyyy"
        ws.cell(row, 9).number_format = "@"
    add_table(ws, "tbl_hr_employee_profile", len(headers))
    add_row_quality_formatting(ws, "K")
    add_list_validation(ws, "F", "dm_gender", "Chọn giới tính chuẩn hóa.")
    add_custom_validation(ws, f"C{FIRST_DATA_ROW}:C{LAST_DATA_ROW}", f'=OR(C{FIRST_DATA_ROW}="",AND(LEN(C{FIRST_DATA_ROW})<=20,ISERROR(SEARCH(" ",C{FIRST_DATA_ROW}))))', "Mã nhân viên không hợp lệ", "Mã tối đa 20 ký tự và không có khoảng trắng.", "Nhập dạng text, ví dụ 00001.")
    add_custom_validation(ws, f"E{FIRST_DATA_ROW}:E{LAST_DATA_ROW}", f'=OR(E{FIRST_DATA_ROW}="",AND(ISNUMBER(E{FIRST_DATA_ROW}),E{FIRST_DATA_ROW}>=DATE(1940,1,1),E{FIRST_DATA_ROW}<=TODAY()))', "Ngày sinh không hợp lệ", "Ngày sinh phải từ 01/01/1940 đến hôm nay.", "Dùng định dạng dd/mm/yyyy.")
    add_custom_validation(ws, f"G{FIRST_DATA_ROW}:H{LAST_DATA_ROW}", f'=OR(G{FIRST_DATA_ROW}="",AND(ISNUMBER(SEARCH("@",G{FIRST_DATA_ROW})),ISNUMBER(SEARCH(".",G{FIRST_DATA_ROW}))))', "Email không hợp lệ", "Email phải có @ và dấu chấm.", "Ví dụ user@company.com.")
    add_custom_validation(ws, f"I{FIRST_DATA_ROW}:I{LAST_DATA_ROW}", f'=OR(I{FIRST_DATA_ROW}="",AND(LEN(I{FIRST_DATA_ROW})>=9,LEN(I{FIRST_DATA_ROW})<=15,ISNUMBER(--I{FIRST_DATA_ROW})))', "Số điện thoại không hợp lệ", "Chỉ nhập 9–15 chữ số.", "Định dạng text để giữ số 0 đầu.")
    ws["C5"].comment = Comment("Khóa nghiệp vụ của nhân viên. Pipeline giữ nguyên phần hiển thị và dùng khóa chuẩn hóa riêng để đối chiếu 00001 với 1.", "Codex")


def build_assignment_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("02_CONG_VIEC")
    title(ws, "BƯỚC 2 · THÔNG TIN CÔNG VIỆC", "Mã nhân viên phải tồn tại ở Bước 1. Dữ liệu phòng ban/chức danh lấy từ danh mục và có thể mở rộng.")
    headers = [
        "Trạng thái kiểm tra", "Lỗi dữ liệu", "Mã nhân viên", "Trạng thái làm việc", "Ngày vào công ty",
        "Chi nhánh", "Phòng ban", "Bộ phận", "Chức danh (VI)", "Position (EN)", "Ngạch", "Cấp bậc",
        "Mã quản lý trực tiếp", "Lương cơ bản", "Ghi chú",
    ]
    style_headers(ws, headers)
    widths = [18, 52, 16, 20, 17, 14, 28, 28, 34, 34, 12, 12, 20, 18, 36]
    style_input_sheet(ws, {1, 2, 10}, widths)
    for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        ws.cell(row, 1, f'=IF($C{row}="","CHƯA NHẬP",IF($B{row}="","HỢP LỆ","LỖI"))')
        ws.cell(row, 2, (
            f'=IF($C{row}="","",TEXTJOIN(" | ",TRUE,'
            f'IF(COUNTIF(\'01_HO_SO\'!$C${FIRST_DATA_ROW}:$C${LAST_DATA_ROW},$C{row})=0,"Mã chưa có ở Bước 1",""),'
            f'IF(COUNTIF($C${FIRST_DATA_ROW}:$C${LAST_DATA_ROW},$C{row})>1,"Trùng mã trong thông tin công việc",""),'
            f'IF($D{row}="","Thiếu trạng thái làm việc",""),'
            f'IF($E{row}="","Thiếu ngày vào công ty",IF(OR(NOT(ISNUMBER($E{row})),$E{row}>TODAY()),"Ngày vào công ty không hợp lệ","")),'
            f'IF($G{row}="","Thiếu phòng ban",""),'
            f'IF($I{row}="","Thiếu chức danh",""),'
            f'IF(AND($M{row}<>"",COUNTIF(\'01_HO_SO\'!$C${FIRST_DATA_ROW}:$C${LAST_DATA_ROW},$M{row})=0),"Mã quản lý không tồn tại",""),'
            f'IF(AND($M{row}<>"",$M{row}=$C{row}),"Không thể tự quản lý chính mình",""),'
            f'IF(OR($N{row}<0,$N{row}>1000000000),"Lương cơ bản ngoài phạm vi","")))'
        ))
        ws.cell(row, 10, f'=IF($I{row}="","",IFERROR(INDEX(DANH_MUC!$D$4:$D$1000,MATCH($I{row},DANH_MUC!$C$4:$C$1000,0)),""))')
        ws.cell(row, 3).number_format = "@"
        ws.cell(row, 5).number_format = "dd/mm/yyyy"
        ws.cell(row, 13).number_format = "@"
        ws.cell(row, 14).number_format = '#,##0 "₫"'
    add_table(ws, "tbl_hr_employment_assignment", len(headers))
    add_row_quality_formatting(ws, "O")
    add_list_validation(ws, "D", "dm_employment_status", "Chọn trạng thái làm việc.")
    add_list_validation(ws, "F", "dm_branch", "Chọn chi nhánh/văn phòng.")
    add_list_validation(ws, "G", "dm_department", "Chọn phòng ban.")
    add_list_validation(ws, "H", "dm_team", "Chọn bộ phận/nhóm.")
    add_list_validation(ws, "I", "dm_job_vi", "Chọn chức danh tiếng Việt; tên tiếng Anh được điền tự động.")
    add_list_validation(ws, "K", "dm_grade", "Chọn ngạch.")
    add_list_validation(ws, "L", "dm_level", "Chọn cấp bậc.")
    add_custom_validation(ws, f"C{FIRST_DATA_ROW}:C{LAST_DATA_ROW}", f'=OR(C{FIRST_DATA_ROW}="",COUNTIF(\'01_HO_SO\'!$C${FIRST_DATA_ROW}:$C${LAST_DATA_ROW},C{FIRST_DATA_ROW})=1)', "Mã nhân viên chưa tồn tại", "Hãy nhập nhân viên tại Bước 1 trước.", "Mã phải trùng chính xác với Bước 1.")
    add_custom_validation(ws, f"E{FIRST_DATA_ROW}:E{LAST_DATA_ROW}", f'=OR(E{FIRST_DATA_ROW}="",AND(ISNUMBER(E{FIRST_DATA_ROW}),E{FIRST_DATA_ROW}<=TODAY()))', "Ngày vào công ty không hợp lệ", "Ngày vào công ty không được ở tương lai.", "Dùng định dạng dd/mm/yyyy.")
    add_custom_validation(ws, f"M{FIRST_DATA_ROW}:M{LAST_DATA_ROW}", f'=OR(M{FIRST_DATA_ROW}="",AND(M{FIRST_DATA_ROW}<>C{FIRST_DATA_ROW},COUNTIF(\'01_HO_SO\'!$C${FIRST_DATA_ROW}:$C${LAST_DATA_ROW},M{FIRST_DATA_ROW})=1))', "Mã quản lý không hợp lệ", "Quản lý phải có ở Bước 1 và không được trùng chính nhân viên.", "Để trống nếu chưa xác định.")
    salary_validation = DataValidation(type="decimal", operator="between", formula1="0", formula2="1000000000", allow_blank=True)
    salary_validation.errorTitle = "Lương không hợp lệ"
    salary_validation.error = "Lương phải từ 0 đến 1.000.000.000."
    salary_validation.errorStyle = "stop"
    salary_validation.showErrorMessage = True
    ws.add_data_validation(salary_validation)
    salary_validation.add(f"N{FIRST_DATA_ROW}:N{LAST_DATA_ROW}")


def build_quality_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("03_KIEM_TRA")
    title(ws, "BƯỚC 3 · KIỂM TRA CHẤT LƯỢNG", "Chỉ xuất/đưa lên SharePoint khi toàn bộ dữ liệu đang nhập hiển thị hợp lệ.")
    ws.sheet_view.showGridLines = False
    for col, width in {"A": 4, "B": 30, "C": 18, "D": 18, "E": 4, "F": 28, "G": 18, "H": 18}.items():
        ws.column_dimensions[col].width = width
    ws["B5"], ws["C5"], ws["D5"] = "Chỉ số", "Hồ sơ", "Công việc"
    for cell in ws[5][1:4]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center")
    metrics = [
        ("Tổng dòng đã nhập", f'=COUNTIF(\'01_HO_SO\'!$C${FIRST_DATA_ROW}:$C${LAST_DATA_ROW},"<>")', f'=COUNTIF(\'02_CONG_VIEC\'!$C${FIRST_DATA_ROW}:$C${LAST_DATA_ROW},"<>")'),
        ("Dòng hợp lệ", f'=COUNTIF(\'01_HO_SO\'!$A${FIRST_DATA_ROW}:$A${LAST_DATA_ROW},"HỢP LỆ")', f'=COUNTIF(\'02_CONG_VIEC\'!$A${FIRST_DATA_ROW}:$A${LAST_DATA_ROW},"HỢP LỆ")'),
        ("Dòng có lỗi", f'=COUNTIF(\'01_HO_SO\'!$A${FIRST_DATA_ROW}:$A${LAST_DATA_ROW},"LỖI")', f'=COUNTIF(\'02_CONG_VIEC\'!$A${FIRST_DATA_ROW}:$A${LAST_DATA_ROW},"LỖI")'),
        ("Mã trùng", f'=COUNTIF(\'01_HO_SO\'!$B${FIRST_DATA_ROW}:$B${LAST_DATA_ROW},"*Trùng mã*")', f'=COUNTIF(\'02_CONG_VIEC\'!$B${FIRST_DATA_ROW}:$B${LAST_DATA_ROW},"*Trùng mã*")'),
        ("Thiếu trường bắt buộc", f'=COUNTIF(\'01_HO_SO\'!$B${FIRST_DATA_ROW}:$B${LAST_DATA_ROW},"*Thiếu*")', f'=COUNTIF(\'02_CONG_VIEC\'!$B${FIRST_DATA_ROW}:$B${LAST_DATA_ROW},"*Thiếu*")'),
    ]
    for row, (label, profile_formula, assignment_formula) in enumerate(metrics, 6):
        ws.cell(row, 2, label)
        ws.cell(row, 3, profile_formula)
        ws.cell(row, 4, assignment_formula)
        for col in range(2, 5):
            ws.cell(row, col).border = Border(bottom=Side(style="thin", color=GRID))
            ws.cell(row, col).alignment = Alignment(horizontal="center" if col > 2 else "left")
    ws.merge_cells("B13:D14")
    ws["B13"] = '=IF(AND(C8=0,D8=0,C6>0,C6=D6),"SẴN SÀNG ĐƯA LÊN SHAREPOINT","CHƯA SẴN SÀNG · HÃY SỬA LỖI")'
    ws["B13"].font = Font(size=14, bold=True, color=WHITE)
    ws["B13"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws["B13"].fill = PatternFill("solid", fgColor=INDIGO)
    ws.conditional_formatting.add("B13:D14", FormulaRule(formula=['LEFT($B$13,8)="SẴN SÀNG"'], fill=PatternFill("solid", fgColor=GREEN)))
    ws.conditional_formatting.add("B13:D14", FormulaRule(formula=['LEFT($B$13,4)="CHƯA"'], fill=PatternFill("solid", fgColor=RED)))
    ws.merge_cells("B16:D18")
    ws["B16"] = "Tiêu chí sẵn sàng: có ít nhất 1 hồ sơ, số hồ sơ bằng số dòng công việc, và cả hai bước không còn dòng LỖI. Bộ kiểm tra trên HR App vẫn chạy lại trước khi ghi dữ liệu."
    ws["B16"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["B16"].font = Font(color=SLATE, size=10)
    ws["B20"] = "Quay lại Bước 1"
    ws["B20"].hyperlink = "#'01_HO_SO'!A1"
    ws["B20"].style = "Hyperlink"
    ws["C20"] = "Quay lại Bước 2"
    ws["C20"].hyperlink = "#'02_CONG_VIEC'!A1"
    ws["C20"].style = "Hyperlink"

    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Chất lượng dữ liệu"
    chart.y_axis.title = ""
    chart.x_axis.title = "Số dòng"
    data = Reference(ws, min_col=3, max_col=4, min_row=5, max_row=8)
    categories = Reference(ws, min_col=2, min_row=6, max_row=8)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 7
    chart.width = 12
    ws.add_chart(chart, "F5")
    ws.protection.sheet = True
    ws.protection.password = "hr-template"


def add_dynamic_name(wb: Workbook, name: str, column: str) -> None:
    formula = f"'DANH_MUC'!${column}$4:INDEX('DANH_MUC'!${column}:${column},MAX(4,COUNTA('DANH_MUC'!${column}:${column})+1))"
    wb.defined_names.add(DefinedName(name, attr_text=formula))


def build_catalog_sheet(wb: Workbook, refs: dict[str, list[str] | list[tuple[str, str]]]) -> None:
    ws = wb.create_sheet("DANH_MUC")
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:I1")
    ws["A1"] = "DANH MỤC CẤU HÌNH"
    ws["A1"].font = Font(size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells("A2:I2")
    ws["A2"] = "Người quản trị có thể bổ sung giá trị ở cuối từng cột. Không sắp xếp riêng cột Chức danh VI/EN vì đây là một cặp mapping."
    ws["A2"].font = Font(color=SLATE, italic=True, size=9)
    headers = ["Trạng thái làm việc", "Giới tính", "Chức danh (VI)", "Position (EN)", "Phòng ban", "Bộ phận", "Chi nhánh", "Ngạch", "Cấp bậc"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(3, col, header)
        cell.fill = PatternFill("solid", fgColor=INDIGO)
        cell.font = Font(bold=True, color=WHITE)
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
        ws.column_dimensions[cell.column_letter].width = 28 if col in {3, 4, 5, 6} else 20
    values: dict[int, list[str]] = {
        1: ["Đang làm việc", "Thử việc", "Tạm nghỉ", "Đã nghỉ việc"],
        2: ["Nam", "Nữ", "Khác", "Không công bố"],
        3: [pair[0] for pair in refs["jobs"]],
        4: [pair[1] for pair in refs["jobs"]],
        5: list(refs["departments"]),
        6: list(refs["teams"]),
        7: list(refs["branches"]),
        8: list(refs["grades"]),
        9: list(refs["levels"]),
    }
    for col, items in values.items():
        for row, value in enumerate(items, 4):
            ws.cell(row, col, value)
            ws.cell(row, col).fill = PatternFill("solid", fgColor=INPUT_FILL)
            ws.cell(row, col).protection = Protection(locked=False)
    ws.freeze_panes = "A4"
    add_dynamic_name(wb, "dm_employment_status", "A")
    add_dynamic_name(wb, "dm_gender", "B")
    add_dynamic_name(wb, "dm_job_vi", "C")
    add_dynamic_name(wb, "dm_job_en", "D")
    add_dynamic_name(wb, "dm_department", "E")
    add_dynamic_name(wb, "dm_team", "F")
    add_dynamic_name(wb, "dm_branch", "G")
    add_dynamic_name(wb, "dm_grade", "H")
    add_dynamic_name(wb, "dm_level", "I")
    ws.protection.sheet = True
    ws.protection.password = "hr-template"
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False


def build_meta_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("SYS_META")
    rows = [
        ("form_key", "hr_employee_onboarding"),
        ("form_version", "1.0.0"),
        ("generated_on", date.today().isoformat()),
        ("profile_table", "tbl_hr_employee_profile"),
        ("assignment_table", "tbl_hr_employment_assignment"),
        ("profile_business_key", "Mã nhân viên"),
        ("assignment_business_key", "Mã nhân viên"),
        ("status_formula", "HỢP LỆ / LỖI / CHƯA NHẬP"),
        ("pipeline_contract", "Read Excel Tables by stable table name; ignore rows with blank Mã nhân viên"),
    ]
    for row, pair in enumerate(rows, 1):
        ws.cell(row, 1, pair[0])
        ws.cell(row, 2, pair[1])
    ws.sheet_state = "veryHidden"


def build_workbook(output: Path, job_file: Path | None, employee_file: Path | None) -> None:
    refs = read_reference_lists(job_file, employee_file)
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    build_start_sheet(wb)
    build_profile_sheet(wb)
    build_assignment_sheet(wb)
    build_quality_sheet(wb)
    build_catalog_sheet(wb, refs)
    build_meta_sheet(wb)
    wb.active = 0
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--job-titles", type=Path)
    parser.add_argument("--employee-master", type=Path)
    args = parser.parse_args()
    build_workbook(args.output, args.job_titles, args.employee_master)
    print(args.output)


if __name__ == "__main__":
    main()
